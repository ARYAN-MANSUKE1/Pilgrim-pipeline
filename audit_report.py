"""
Language audit of every temple -> one Excel workbook.

Read-only. Works from cached JSON downloads of the temple list (id, slug,
title, link, acf) so the live site is only hit once.

Usage:
    python audit_report.py live_cache.json [staging_cache.json] out.xlsx

Sheets: Summary | Temples (one row per temple, one column per language) |
Problem fields (one row per bad field) | Not translated.
"""

from __future__ import annotations

import collections
import json
import sys
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
from app.tasks import wp_client as wp  # noqa: E402
from app.tasks.field_checks import (MIN_CHARS, TOKEN_RE, _mr_hi, content_problem, letter_counts,  # noqa: E402
                                    mr_hi_duplicate, script_frac, visible)
import re  # noqa: E402

LANGS = ["mr", "en", "hi", "gj", "ta", "te", "ml", "kn", "bn", "pa"]
NAME = {"mr": "Marathi", "en": "English", "hi": "Hindi", "gj": "Gujarati", "ta": "Tamil", "te": "Telugu",
        "ml": "Malayalam", "kn": "Kannada", "bn": "Bengali", "pa": "Punjabi"}
BUCKET_OF = {"mr": "deva", "hi": "deva", "en": "latin", "gj": "gj", "ta": "ta", "te": "te", "ml": "ml",
             "kn": "kn", "bn": "bn", "pa": "pa"}
CODES_OF = {"deva": ["mr", "hi"], "latin": ["en"], "gj": ["gj"], "ta": ["ta"], "te": ["te"], "ml": ["ml"],
            "kn": ["kn"], "bn": ["bn"], "pa": ["pa"]}
RANGES = [("deva", 0x0900, 0x097F), ("bn", 0x0980, 0x09FF), ("pa", 0x0A00, 0x0A7F), ("gj", 0x0A80, 0x0AFF),
          ("ta", 0x0B80, 0x0BFF), ("te", 0x0C00, 0x0C7F), ("kn", 0x0C80, 0x0CFF), ("ml", 0x0D00, 0x0D7F)]


def dominant(text: str) -> tuple[str, float]:
    """(script bucket holding most letters, its share of all letters)."""
    n = collections.Counter()
    for c, k in letter_counts(text).items():  # distinct characters only - fast on long articles
        if c.isascii():
            n["latin"] += k
            continue
        for b, lo, hi in RANGES:
            if lo <= ord(c) <= hi:
                n[b] += k
                break
        else:
            n["other"] += k
    total = sum(n.values())
    if not total:
        return "none", 0.0
    b, k = n.most_common(1)[0]
    return b, k / total


def label(bucket: str, text: str) -> str:
    if bucket == "latin":
        return "English"
    if bucket == "deva":
        mr_c, hi_c = _mr_hi(text)
        return "Hindi" if hi_c > mr_c else "Marathi" if mr_c > hi_c else "Marathi/Hindi"
    return NAME.get(bucket, "unknown-script")


def norm(s: str) -> str:
    return re.sub(r"\W+", "", s)


GROUP_OF = {"Untranslated": "Untranslated (source text left in the field)",
            "Wrong language": "Wrong language (another language's text in the field)",
            "Mixed": "Mixed / partly translated",
            "Junk marker": "Junk marker text (zZqSEGqZz)",
            "Marathi/Hindi": "Marathi and Hindi mixed up",
            "Missing": "Missing (empty)",
            "Title": "Title problem"}


def analyze(acf: dict) -> dict:
    """{lang: {'content': info, 'title': info}} where info has status/cat/group/detail/actual/rel/excerpt."""
    vis = {l: visible(acf.get(wp.CONTENT_FIELD_KEYS[l]) or "") for l in LANGS}
    dom = {l: dominant(vis[l])[0] for l in LANGS}
    nrm = {l: norm(vis[l]) for l in LANGS}
    out = {}
    for l in LANGS:
        raw = acf.get(wp.CONTENT_FIELD_KEYS[l]) or ""
        v = vis[l]
        info = dict(status="OK", cat="", group="", detail="", actual="", rel="", excerpt=v[:90])
        if len(v) < MIN_CHARS:
            info.update(status="MISSING", cat="Missing (empty)", group="Missing", detail=f"{len(v)} visible characters", excerpt="")
        else:
            d, share = dominant(v)
            sf = script_frac(v, l)
            cat = group = ""
            if sf < (0.85 if l == "en" else 0.6):
                actual = label(d, v)
                info["actual"] = actual
                if share >= 0.6 and d == "latin":
                    cat, group = "Untranslated - English text left in the field", "Untranslated"
                elif share >= 0.6 and d == "deva":
                    cat, group = f"Untranslated - {actual} text left in the field", "Untranslated"
                elif share >= 0.6:
                    cat, group = f"Wrong language - {actual} text in the {NAME[l]} field", "Wrong language"
                else:
                    cat, group = f"Mixed languages / partly translated ({sf:.0%} in the right script)", "Mixed"
                rels = []
                for c in CODES_OF.get(d, []):
                    if c != l and nrm[c] and nrm[c] == nrm[l]:
                        rels.append(f"exact copy of the {NAME[c]} field")
                    if c != l and dom[c] == BUCKET_OF[l] and d != BUCKET_OF[l] and BUCKET_OF[c] == d:
                        rels.append(f"SWAPPED with the {NAME[c]} field")
                info["rel"] = "; ".join(rels)
            elif l in ("mr", "hi") and mr_hi_duplicate(acf.get(wp.CONTENT_FIELD_KEYS["mr"]) or "",
                                                       acf.get(wp.CONTENT_FIELD_KEYS["hi"]) or "") == l:
                cat, group = f"Marathi and Hindi fields hold identical text (this one is the wrong-language copy)", "Marathi/Hindi"
            else:
                reason = content_problem(raw, l)
                if reason == "token leak":
                    cat, group = "Junk marker text (zZqSEGqZz) inside the article", "Junk marker"
                elif reason in ("Hindi text in Marathi field", "Marathi text in Hindi field"):
                    cat, group = reason, "Marathi/Hindi"
                elif reason == "mostly Latin":
                    cat, group = "Mixed - English words left inside the text", "Mixed"
            if cat and TOKEN_RE.search(raw) and group != "Junk marker":
                info["detail"] = "also contains the junk marker zZqSEGqZz"
            if cat:
                info.update(status="PROBLEM", cat=cat, group=group)
        t = (acf.get(wp.TITLE_FIELD_KEYS[l]) or "").strip()
        tinfo = dict(status="OK", cat="", group="", detail="", actual="", rel="", excerpt=t[:90])
        if l != "en":
            has = len(v) >= MIN_CHARS
            if not t and has:
                tinfo.update(status="PROBLEM", cat="Title missing", group="Title")
            elif t and TOKEN_RE.search(t):
                tinfo.update(status="PROBLEM", cat="Title contains junk marker text", group="Title")
            elif t and script_frac(t, l) < 0.5:
                d, _ = dominant(t)
                a = label(d, t)
                tinfo.update(status="PROBLEM", group="Title", actual=a,
                             cat=f"Title in wrong language - {a} title in the {NAME[l]} title field")
        out[l] = {"content": info, "title": tinfo}
    return out


def short(info: dict) -> str:
    if info["status"] == "OK":
        return "OK"
    if info["status"] == "MISSING":
        return "MISSING"
    g = info["group"]
    s = {"Untranslated": f"UNTRANSLATED ({info['actual']})", "Wrong language": f"WRONG LANGUAGE ({info['actual']})",
         "Mixed": "MIXED", "Junk marker": "JUNK TEXT", "Marathi/Hindi": "MR/HI MIX-UP"}[g]
    return s + (" - SWAPPED" if "SWAPPED" in info["rel"] else "")


FILL = {"OK": "E2F0D9", "MISSING": "D9D9D9", "PROBLEM": "F8CBAD", "MIXED": "FFE699"}
HEAD = PatternFill("solid", fgColor="1F3864")


def build(live: list, stage: list | None, out_path: str) -> dict:
    stage_by_slug = {t["slug"]: t for t in (stage or [])}
    stage_an = {s: analyze(t.get("acf") or {}) for s, t in stage_by_slug.items()} if stage else {}
    temples_rows, problem_rows, untranslated_rows = [], [], []
    matrix = collections.defaultdict(collections.Counter)  # group -> lang -> n
    title_by_lang = collections.Counter()
    rel_counter = collections.Counter()
    tstat = collections.Counter()

    for t in sorted(live, key=lambda x: x["id"]):
        acf = t.get("acf") or {}
        an = analyze(acf)
        name = (t.get("title") or {}).get("raw") or (t.get("title") or {}).get("rendered") or t["slug"]
        link = t.get("link", "")
        n_missing = sum(an[l]["content"]["status"] == "MISSING" for l in LANGS)
        probs = [(l, k, an[l][k]) for l in LANGS for k in ("content", "title") if an[l][k]["status"] == "PROBLEM"]
        ok_langs = [NAME[l] for l in LANGS if an[l]["content"]["status"] == "OK" and an[l]["title"]["status"] == "OK"]
        n_audio = sum(bool(acf.get(wp.AUDIO_FIELD_KEYS[l])) for l in LANGS)
        source_ok = any(an[l]["content"]["status"] == "OK" for l in ("en", "mr", "hi"))
        if n_missing >= 8:
            overall, fix = "NOT TRANSLATED", "Full translation needed (no translations exist yet)"
            untranslated_rows.append([t["id"], name, link, 10 - n_missing, n_audio])
        elif probs or n_missing:
            overall = "NEEDS FIX"
            fix = (f"Retranslate {len(probs)} field(s) with repair_fields.py" if source_ok else "No clean source language - manual")
            if n_missing:
                fix += f"; {n_missing} language(s) empty"
        else:
            overall, fix = "OK", "None"
        tstat[overall] += 1

        s_note = "n/a"
        if stage is not None:
            if t["slug"] not in stage_by_slug:
                s_note = "Not on staging (live-only temple)"
            elif probs:
                sa = stage_an[t["slug"]]
                same = sum(sa[l][k]["status"] == "PROBLEM" for l, k, _ in probs)
                s_note = ("Yes - same problem on staging (source issue)" if same == len(probs)
                          else "Live only (staging is fine)" if same == 0 else f"Partly ({same} of {len(probs)} also on staging)")
        summary_bits = []
        for l, k, info in probs:
            what = info["cat"] + (f" [{info['rel']}]" if info["rel"] else "")
            summary_bits.append(f"{NAME[l]}{' title' if k == 'title' else ''}: {what}")
            matrix[info["group"]][l] += 1
            if k == "title":
                title_by_lang[l] += 1
            if k == "content" and info["group"] in ("Untranslated", "Wrong language", "Mixed"):
                kind = ("SWAPPED with another field" if "SWAPPED" in info["rel"] else
                        "exact copy of another field" if info["rel"] else "not a copy of any other field")
                rel_counter[f"{GROUP_OF[info['group']]}: {kind}"] += 1
            problem_rows.append([t["id"], name, link, NAME[l], "Title" if k == "title" else "Article text",
                                 GROUP_OF[info["group"]], info["cat"], info["actual"], info["rel"],
                                 info["detail"], info["excerpt"],
                                 (("Yes" if stage_an.get(t["slug"], {}).get(l, {}).get(k, {}).get("status") == "PROBLEM" else "No")
                                  if t["slug"] in stage_by_slug else "Not on staging" if stage is not None else "n/a"),
                                 "Retranslate from " + ("English" if an["en"]["content"]["status"] == "OK" else
                                                        "Marathi" if an["mr"]["content"]["status"] == "OK" else "Hindi")
                                 if source_ok else "No clean source language"])
        for l in LANGS:
            if an[l]["content"]["status"] == "MISSING" and n_missing < 8:
                matrix["Missing"][l] += 1
                problem_rows.append([t["id"], name, link, NAME[l], "Article text", GROUP_OF["Missing"], "Missing (empty)",
                                     "", "", an[l]["content"]["detail"], "", "n/a", "Needs translation"])
        cells = []
        for l in LANGS:
            c, ti = an[l]["content"], an[l]["title"]
            txt = short(c) + (" + title wrong" if ti["status"] == "PROBLEM" else "")
            cells.append((txt, "OK" if txt == "OK" else "MISSING" if txt == "MISSING" else "PROBLEM"))
        temples_rows.append(([t["id"], name, link, overall, f"{len(ok_langs)}/10", ", ".join(ok_langs) or "none",
                              len(probs) + (n_missing if n_missing < 8 else 0), "; ".join(summary_bits) or "",
                              *[c[0] for c in cells], n_audio, s_note, fix], [c[1] for c in cells], overall))

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Language audit of the live site"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {time.strftime('%Y-%m-%d %H:%M')}. Every temple, all 10 languages. Read-only check; nothing was changed."])
    ws.append([])
    ws.append(["Temples checked", len(live)])
    for k in ("OK", "NEEDS FIX", "NOT TRANSLATED"):
        ws.append([f"Temples: {k}", tstat[k]])
    ws.append([])
    ws.append(["Problem type (article text unless noted)"] + [NAME[l] for l in LANGS] + ["Total"])
    hdr = ws.max_row
    for g in ["Untranslated", "Wrong language", "Mixed", "Junk marker", "Marathi/Hindi", "Missing", "Title"]:
        row = [GROUP_OF[g]] + [matrix[g][l] for l in LANGS]
        ws.append(row + [sum(row[1:])])
    ws.append([])
    ws.append(["Are the wrong fields copies / swaps of each other?"])
    for k, v in sorted(rel_counter.items()):
        ws.append([f"   {k}", v])
    ws.append([])
    ws.append(["How to read it"])
    for line in ["UNTRANSLATED: the field holds the source text (English or Marathi/Hindi) - the translation step never wrote its own language.",
                 "WRONG LANGUAGE: the field holds a different Indian language's text (e.g. Bengali text in the Kannada field). 'SWAPPED' means two fields hold each other's text.",
                 "MIXED: part of the text is in the right script, part is not.",
                 "JUNK TEXT: the batching marker zZqSEGqZz leaked into the article.",
                 "Sheet 'Temples' has one row per temple; 'Problem fields' has one row per broken field; 'Not translated' lists temples with no translations."]:
        ws.append([line])
    for c in ws[hdr]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    ws.column_dimensions["A"].width = 62
    for i in range(2, 13):
        ws.column_dimensions[get_column_letter(i)].width = 12

    def sheet(title, header, rows, widths, colors=None, link_col=3):
        w = wb.create_sheet(title)
        w.append(header)
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for ri, r in enumerate(rows, 2):
            w.append(r)
            lc = w.cell(row=ri, column=link_col)
            if isinstance(lc.value, str) and lc.value.startswith("http"):
                lc.hyperlink = lc.value
                lc.font = Font(color="0563C1", underline="single")
            for cc in w[ri]:
                cc.alignment = Alignment(wrap_text=True, vertical="top")
            if colors:
                colors(w, ri)
        for i, wd in enumerate(widths, 1):
            w.column_dimensions[get_column_letter(i)].width = wd
        w.freeze_panes = "C2"
        w.auto_filter.ref = w.dimensions
        return w

    th = ["Temple ID", "Temple name", "Link", "Overall", "Languages OK", "Which languages are fine", "Problems",
          "What is wrong"] + [NAME[l] for l in LANGS] + ["Audio (of 10)", "Same problem on staging?", "Fix needed"]
    tw = sheet("Temples", th, [r[0] for r in temples_rows], [9, 38, 30, 16, 9, 30, 9, 70] + [22] * 10 + [9, 28, 34])
    for ri, (row, stats, overall) in enumerate(temples_rows, 2):
        tw.cell(row=ri, column=4).fill = PatternFill("solid", fgColor={"OK": FILL["OK"], "NEEDS FIX": FILL["PROBLEM"], "NOT TRANSLATED": "BFBFBF"}[overall])
        for j, s in enumerate(stats):
            tw.cell(row=ri, column=9 + j).fill = PatternFill("solid", fgColor=FILL[s])
    sheet("Problem fields", ["Temple ID", "Temple name", "Link", "Language", "Field", "Problem type", "What is wrong",
                             "What language the text really is", "Swapped / copied?", "Extra", "Start of the text",
                             "Same problem on staging?", "Fix"], problem_rows,
          [9, 34, 30, 12, 12, 34, 52, 16, 30, 26, 44, 16, 26])
    sheet("Not translated", ["Temple ID", "Temple name", "Link", "Languages with text (of 10)", "Audio (of 10)"],
          untranslated_rows, [9, 44, 34, 14, 12])
    wb.save(out_path)
    return {"temples": len(live), "status": dict(tstat), "problem_fields": len(problem_rows),
            "by_group": {g: dict(c) for g, c in matrix.items()}, "wrong_language_relation": dict(rel_counter),
            "titles_by_lang": dict(title_by_lang)}


if __name__ == "__main__":
    args = sys.argv[1:]
    live = json.load(open(args[0], encoding="utf-8"))
    stage = json.load(open(args[1], encoding="utf-8")) if len(args) == 3 else None
    print(json.dumps(build(live, stage, args[-1]), ensure_ascii=False, indent=1))
